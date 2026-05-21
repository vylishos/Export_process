import streamlit as st
import pandas as pd
import io
from openpyxl.styles import Font, Border, Side, Alignment
from datetime import datetime

st.set_page_config(page_title="Export processor", layout="wide")

def apply_sklad_formatting(writer, df_sorted):
    worksheet = writer.sheets['Sklad']
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    thick_bottom_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thick'))
    header_font = Font(bold=True)
    date_font = Font(bold=True, size=11)

    worksheet['A1'].font = date_font

    for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=4)):
        for cell in row:
            if row_idx == 0:
                cell.font = header_font
                cell.border = thick_bottom_border
            else:
                cell.border = thin_border

    if not df_sorted.empty:
        max_length = max([len(str(val)) for val in df_sorted['Název']] + [10])
        worksheet.column_dimensions['A'].width = max_length + 2

    for col in ['B', 'C', 'D']:
        worksheet.column_dimensions[col].width = 15

def apply_prehled_formatting(writer, df_vystup):
    worksheet = writer.sheets['Prehled']
    max_row = worksheet.max_row
    max_col = worksheet.max_column
    thin_side = Side(style='thin')
    thick_side = Side(style='thick')
    top_border_only = Border(top=thin_side)
    header_border = Border(bottom=thick_side)
    header_font = Font(bold=True)
    date_font = Font(bold=True, size=11)
    no_border = Border()
    align_right = Alignment(horizontal='right')

    worksheet['A1'].font = date_font

    for r_idx in range(2, max_row + 1):
        for c_idx in range(1, max_col + 1):
            cell = worksheet.cell(row=r_idx, column=c_idx)
            
            if c_idx == 1:
                cell.alignment = align_right

            if r_idx == 2:
                cell.font = header_font
                cell.border = header_border
            else:
                val_a = worksheet.cell(row=r_idx, column=1).value
                is_new_order = val_a is not None and str(val_a).strip() != "" and str(val_a).strip() not in ["Zásilkovna", "GLS", "UPS"]
                cell.border = top_border_only if is_new_order else no_border

    for i, col_name in enumerate(df_vystup.columns):
        column_data = df_vystup[col_name].astype(str).fillna('')
        clean_lengths = column_data.replace(['nan', 'None'], '').apply(len)
        
        max_content_len = clean_lengths.max() if not clean_lengths.empty else 0
        max_len = max(max_content_len, len(str(col_name))) + 2
        
        col_letter = chr(65 + i) if i < 26 else f"A{chr(65 + (i-26))}"
        worksheet.column_dimensions[col_letter].width = max_len

st.title("Export processor")
uploaded_file = st.file_uploader("Nahrajte exportní Excel soubor (.xlsx)", type=['xlsx'])

if uploaded_file:
    df = pd.read_excel(uploaded_file, engine='openpyxl')
    st.success("Soubor úspěšně nahrán!")

    datum_soubor = datetime.now().strftime("%d.%m.")
    datum_text = datetime.now().strftime("%d.%m.%Y")

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("1. Skladový seznam")
        try:
            df1 = df.iloc[:, [25, 28, 26, 30]].copy()
            df1.columns = ['Název', 'Reference', 'Varianta', 'Ks']
            
            df1 = df1.dropna(subset=['Reference']).copy()
            df1['Ks'] = pd.to_numeric(df1['Ks'], errors='coerce').fillna(0).astype(int)
            
            df1['Varianta'] = df1['Varianta'].astype(str).replace('nan', '')
            df1 = df1[df1['Ks'] > 0].sort_values(by=['Varianta'])
            
            output1 = io.BytesIO()
            with pd.ExcelWriter(output1, engine='openpyxl') as writer:
                df1.to_excel(writer, index=False, sheet_name='Sklad', startrow=1)
                worksheet = writer.sheets['Sklad']
                worksheet['A1'] = f"Export ze dne: {datum_text}"
                apply_sklad_formatting(writer, df1)
            
            st.download_button(
                label="Stáhnout Skladový seznam", 
                data=output1.getvalue(), 
                file_name=f"seznam_pro_sklad_{datum_soubor}xlsx"
            )
        except Exception as e:
            st.error(f"Chyba při tvorbě skladu: {e}")

    with col2:
        st.subheader("2. Přehled objednávek")
        try:
            is_main_order = df.iloc[:, 0].notna()
            is_item = df.iloc[:, 28].notna() & (df.iloc[:, 28].astype(str).str.strip() != '')
            is_shipping = is_item.shift(1, fill_value=False) & ~is_item
            
            # Příprava sběrných seznamů pro jednotlivé přehledy
            vystup_rows = []
            zasilkovna_rows = []
            gls_rows = []
            ups_rows = []
            
            current_order_rows = []
            
            for idx, row in df.iterrows():
                if is_main_order[idx]:
                    current_order_rows = [{
                        'Číslo objednávky': row.iloc[0],
                        'Jméno': row.iloc[2],
                        'Reference': None,
                        'Varianta': None,
                        'Ks': None
                    }]
                elif is_item[idx]:
                    current_order_rows.append({
                        'Číslo objednávky': row.iloc[0],
                        'Jméno': row.iloc[2],
                        'Reference': row.iloc[28],
                        'Varianta': row.iloc[26],
                        'Ks': row.iloc[30]
                    })
                elif is_shipping[idx]:
                    puvodni_text = str(row.iloc[25]) if pd.notna(row.iloc[25]) else ""
                    
                    if "UPS" in puvodni_text:
                        cisty_dopravce = "UPS"
                    elif "GLS" in puvodni_text:
                        cisty_dopravce = "GLS"
                    else:
                        cisty_dopravce = "Zásilkovna"
                    
                    current_order_rows.append({
                        'Číslo objednávky': cisty_dopravce,
                        'Jméno': None,
                        'Reference': None,
                        'Varianta': None,
                        'Ks': None
                    })
                    current_order_rows.append({
                        'Číslo objednávky': None,
                        'Jméno': None,
                        'Reference': None,
                        'Varianta': None,
                        'Ks': None
                    })
                    
                    # Rozdistribuování nasbírané objednávky do správných seznamů
                    vystup_rows.extend(current_order_rows)
                    
                    if cisty_dopravce == "UPS":
                        ups_rows.extend(current_order_rows)
                    elif cisty_dopravce == "GLS":
                        gls_rows.extend(current_order_rows)
                    else:
                        zasilkovna_rows.extend(current_order_rows)
                        
                    current_order_rows = []
            
            # 1. GENEROVÁNÍ KOMPLETNÍHO PŘEHLEDU
            df_vystup = pd.DataFrame(vystup_rows, columns=['Číslo objednávky', 'Jméno', 'Reference', 'Varianta', 'Ks'])
            output2 = io.BytesIO()
            with pd.ExcelWriter(output2, engine='openpyxl') as writer:
                df_vystup.to_excel(writer, index=False, sheet_name='Prehled', startrow=1)
                worksheet = writer.sheets['Prehled']
                worksheet['A1'] = f"Export ze dne: {datum_text}"
                apply_prehled_formatting(writer, df_vystup)
            
            st.download_button(
                label="Stáhnout Kompletní přehled", 
                data=output2.getvalue(), 
                file_name=f"Prehled_objednavek_{datum_soubor}xlsx"
            )
            
            # SEKCE PRO JEDNOTLIVÉ DOPRAVCE
            st.markdown("---")
            st.markdown("### 📦 Objednávky podle dopravců")
            
            # Generování Zásilkovny
            if zasilkovna_rows:
                df_zasilkovna = pd.DataFrame(zasilkovna_rows, columns=['Číslo objednávky', 'Jméno', 'Reference', 'Varianta', 'Ks'])
                output_z = io.BytesIO()
                with pd.ExcelWriter(output_z, engine='openpyxl') as writer:
                    df_zasilkovna.to_excel(writer, index=False, sheet_name='Prehled', startrow=1)
                    worksheet = writer.sheets['Prehled']
                    worksheet['A1'] = f"Export Zásilkovna ze dne: {datum_text}"
                    apply_prehled_formatting(writer, df_zasilkovna)
                st.download_button(
                    label="Stáhnout Přehled - Zásilkovna", 
                    data=output_z.getvalue(), 
                    file_name=f"Prehled_Zasilkovna_{datum_soubor}xlsx"
                )
                
            # Generování GLS
            if gls_rows:
                df_gls = pd.DataFrame(gls_rows, columns=['Číslo objednávky', 'Jméno', 'Reference', 'Varianta', 'Ks'])
                output_g = io.BytesIO()
                with pd.ExcelWriter(output_g, engine='openpyxl') as writer:
                    df_gls.to_excel(writer, index=False, sheet_name='Prehled', startrow=1)
                    worksheet = writer.sheets['Prehled']
                    worksheet['A1'] = f"Export GLS ze dne: {datum_text}"
                    apply_prehled_formatting(writer, df_gls)
                st.download_button(
                    label="Stáhnout Přehled - GLS", 
                    data=output_g.getvalue(), 
                    file_name=f"Prehled_GLS_{datum_soubor}xlsx"
                )
                
            # Generování UPS
            if ups_rows:
                df_ups = pd.DataFrame(ups_rows, columns=['Číslo objednávky', 'Jméno', 'Reference', 'Varianta', 'Ks'])
                output_u = io.BytesIO()
                with pd.ExcelWriter(output_u, engine='openpyxl') as writer:
                    df_ups.to_excel(writer, index=False, sheet_name='Prehled', startrow=1)
                    worksheet = writer.sheets['Prehled']
                    worksheet['A1'] = f"Export UPS ze dne: {datum_text}"
                    apply_prehled_formatting(writer, df_ups)
                st.download_button(
                    label="Stáhnout Přehled - UPS", 
                    data=output_u.getvalue(), 
                    file_name=f"Prehled_UPS_{datum_soubor}xlsx"
                )
                
        except Exception as e:
            st.error(f"Chyba při tvorbě přehledů: {e}")